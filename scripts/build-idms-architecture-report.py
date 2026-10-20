"""Build the architecture assessment deliverable; does not change platform state."""
from pathlib import Path
import re
import zipfile
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/idms-architecture-gap-business-2026-10-15.md'
OUTPUT = ROOT / 'outputs/IDMS多设备数据管理_架构缺陷及盈利模式_2026-10-15_v1.2.docx'

def shade(cell, fill):
    item = OxmlElement('w:shd')
    item.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(item)

def font(run, size=10.5, bold=False, color=None):
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def text(paragraph, value, size=10.5, bold=False, color=None):
    # The report uses Markdown only for structure; keep content editable in Word.
    value = value.replace('`', '')
    font(paragraph.add_run(value), size, bold, color)
    return paragraph

def set_repeat(row):
    header = OxmlElement('w:tblHeader')
    header.set(qn('w:val'), 'true')
    row._tr.get_or_add_trPr().append(header)

def no_split(row):
    node = OxmlElement('w:cantSplit')
    row._tr.get_or_add_trPr().append(node)

def make_table(doc, rows):
    width = 17.0 / len(rows[0])
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = 'Table Grid'
    table.autofit = False
    for i, value in enumerate(rows[0]):
        cell = table.rows[0].cells[i]
        cell.width = Cm(width)
        shade(cell, '183D4D')
        text(cell.paragraphs[0], value, 9, True, 'FFFFFF')
    set_repeat(table.rows[0])
    for index, values in enumerate(rows[1:]):
        row = table.add_row()
        no_split(row)
        for i, cell in enumerate(row.cells):
            cell.width = Cm(width)
            if index % 2 == 0:
                shade(cell, 'F1F5F7')
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.08
            text(paragraph, values[i] if i < len(values) else '', 9)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def build():
    source = SOURCE.read_text(encoding='utf-8')
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
    normal.paragraph_format.line_spacing = 1.22
    normal.paragraph_format.space_after = Pt(6)
    for name in ['Heading 1', 'Heading 2', 'Heading 3']:
        style = doc.styles[name]
        style.font.color.rgb = RGBColor.from_string('183D4D')
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
        style.paragraph_format.keep_with_next = True
    text(section.header.paragraphs[0], 'IDMS  |  产品架构与商业化评审稿', 8, color='607D8B')
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    text(footer, '评审稿 · 非生产验收证明  |  ', 8, color='607D8B')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.title = 'IDMS 多设备数据管理：架构、缺陷与盈利模式'
    doc.core_properties.subject = '可复用产品架构、商用差距、定价实验与试点验收'
    doc.core_properties.author = 'IDMS 产品方案'

    cover = doc.add_paragraph()
    cover.paragraph_format.space_before = Pt(64)
    text(cover, 'IDMS', 38, True, '183D4D')
    text(doc.add_paragraph(), '多类型设备智能数据管理\n与维护决策平台', 25, True, '183D4D')
    text(doc.add_paragraph(), '整体架构 · 功能缺陷 · 盈利模式', 16, color='367B8C')
    text(doc.add_paragraph(), '评审稿 1.2  |  2026-10-15', 11, color='607D8B')
    text(doc.add_paragraph(), '决策摘要：多类型设备，先做数据管理。以统一台账、测点、上传标准化和可信追溯形成基础版，预测维护按设备族验证后增购。', 12)
    text(doc.add_paragraph(), '范围声明：基于现有代码与集成证据的设计评估。价格与收益仅为测算假设；不代表生产商用验收。本次没有部署未验收功能。', 10, color='806424')
    doc.add_page_break()
    text(doc.add_paragraph(), '阅读导览', 18, True, '183D4D')
    for line in ['一、结论与产品边界', '二、总体系统架构', '三、数据架构与追溯设计', '四、预测算法与 AI 工作流', '五、当前系统功能缺陷与风险优先级', '六、模板、插件与接口产品化', '七、盈利模式与建议定价', '八、阶段实施与验收路径', '九、决策建议', '附录：主要评估依据']:
        text(doc.add_paragraph(), line, 11)
    doc.add_page_break()

    lines = source.splitlines()
    index = 0
    headings = []
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith('# '):
            index += 1
            continue
        if line.startswith('## 整体架构'):
            index += 1
            continue
        if line.startswith('```'):
            index += 1
            block = []
            while index < len(lines) and not lines[index].startswith('```'):
                block.append(lines[index])
                index += 1
            table = doc.add_table(rows=1, cols=1)
            shade(table.cell(0, 0), 'EAF2F5')
            paragraph = table.cell(0, 0).paragraphs[0]
            paragraph.alignment = 1
            text(paragraph, '\n'.join(block), 9, color='183D4D')
            index += 1
            continue
        if line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].startswith('|'):
                cells = [value.strip() for value in lines[index].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', value.replace(' ', '')) for value in cells):
                    rows.append(cells)
                index += 1
            make_table(doc, rows)
            continue
        if line.startswith('## '):
            heading = line[3:]
            headings.append(heading)
            paragraph = doc.add_paragraph(style='Heading 1')
            text(paragraph, heading, 17, True, '183D4D')
        elif line.startswith('### '):
            paragraph = doc.add_paragraph(style='Heading 2')
            text(paragraph, line[4:], 12, True, '183D4D')
        elif line.startswith('- '):
            text(doc.add_paragraph(style='List Bullet'), line[2:])
        elif re.match(r'^\d+\. ', line):
            text(doc.add_paragraph(style='List Number'), re.sub(r'^\d+\. ', '', line))
        else:
            text(doc.add_paragraph(), line)
        index += 1
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    # Verify a well-formed OOXML package, and the substantive content and tables.
    with zipfile.ZipFile(OUTPUT) as package:
        assert package.testzip() is None
        assert 'word/document.xml' in package.namelist()
    saved = Document(OUTPUT)
    body = '\n'.join(p.text for p in saved.paragraphs)
    assert len(headings) == 10, headings
    assert len(saved.tables) == 9
    for required in ['P0', 'ROI', '8.6', 'SHADOW', '不是生产安全审计', '40 小时', 'RUL', '多类型设备，先做数据管理', '53.3%', '41.2%', 'typed-sql-read-model', 'ISOLATED']:
        all_text = body + '\n'.join(c.text for t in saved.tables for r in t.rows for c in r.cells)
        assert required in all_text, required
    print(f'Created: {OUTPUT.name}')
    print(f'Sections: {len(headings)}; tables: {len(saved.tables)}; bytes: {OUTPUT.stat().st_size}')
    print('OOXML integrity and required content verified; production acceptance: false')

if __name__ == '__main__':
    build()
